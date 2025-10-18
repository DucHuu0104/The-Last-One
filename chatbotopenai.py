import streamlit as st
from langchain_community.document_loaders import PyPDFLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_community.vectorstores import FAISS
from langchain.chains.question_answering import load_qa_chain
from langchain.prompts import PromptTemplate
from docx import Document
from dotenv import load_dotenv
from pptx import Presentation
import tempfile
import os
from io import BytesIO
import matplotlib.pyplot as plt
from collections import Counter
import seaborn as sns
from wordcloud import STOPWORDS
from PIL import Image
import pytesseract

# =============================
# 🔐 Load biến môi trường
# =============================
load_dotenv()
api_key = os.getenv("OPENAI_API_KEY")

if not api_key:
    st.error("❌ Không tìm thấy API key. Kiểm tra lại file .env")
    st.stop()

# =============================
# 📚 Hàm xử lý tài liệu (PDF, DOCX, PPTX, TXT, hình ảnh)
# =============================
def get_text(docs):
    text = ""
    try:
        for file in docs:
            if file.name.endswith(".pdf"):
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                    tmp.write(file.read())
                    pdf_path = tmp.name
                pdf_reader = PyPDFLoader(pdf_path)
                for page in pdf_reader.load_and_split():
                    text += page.page_content
                os.unlink(pdf_path)

            elif file.name.endswith(".docx"):
                with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                    tmp.write(file.read())
                    doc_path = tmp.name
                doc = Document(doc_path)
                for para in doc.paragraphs:
                    text += para.text + "\n"
                os.unlink(doc_path)

            elif file.name.endswith(".pptx"):
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
                    tmp.write(file.read())
                    ppt_path = tmp.name
                prs = Presentation(ppt_path)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            text += shape.text + "\n"
                os.unlink(ppt_path)

            elif file.name.endswith(".txt"):
                text += file.read().decode("utf-8") + "\n"

            elif file.name.endswith((".png", ".jpg", ".jpeg", ".bmp")):
                with tempfile.NamedTemporaryFile(delete=False, suffix=os.path.splitext(file.name)[1]) as tmp:
                    tmp.write(file.read())
                    img_path = tmp.name
                try:
                    img = Image.open(img_path)
                    text += pytesseract.image_to_string(img, lang="vie+eng") + "\n"
                except Exception as e:
                    st.error(f"Lỗi OCR với {file.name}: {e}")
                finally:
                    os.unlink(img_path)
            else:
                st.warning(f"⚠️ Không hỗ trợ định dạng: {file.name}")
    except Exception as e:
        st.error(f"Lỗi khi đọc tài liệu: {e}")
    return text


# =============================
# ✂️ Chia nhỏ văn bản
# =============================
def get_text_chunk(text):
    try:
        splitter = RecursiveCharacterTextSplitter(chunk_size=10000, chunk_overlap=1000)
        chunks = splitter.split_text(text)
        return chunks
    except Exception as e:
        st.error(f"Lỗi chia chunk: {e}")
        return []


# =============================
# 💾 Tạo và lưu FAISS vector database
# =============================
def get_vector_store(text_chunks):
    try:
        embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
        vector_store = FAISS.from_texts(text_chunks, embedding=embeddings)
        vector_store.save_local("faiss_index")
        st.success("✅ Tài liệu đã được phân tích và lưu xong.")
    except Exception as e:
        st.error(f"Lỗi lưu FAISS: {e}")


# =============================
# 🧠 Tạo chuỗi hỏi đáp (QA Chain)
# =============================
def get_conversation_chain():
    prompt_template = """
    Trả lời câu hỏi một cách chi tiết nhất có thể dựa trên ngữ cảnh cung cấp.
    Nếu câu trả lời không nằm trong ngữ cảnh, hãy nói "Câu trả lời không có trong ngữ cảnh".

    Ngữ cảnh: {context}
    Câu hỏi: {question}

    Trả lời:
    """
    try:
        model = ChatOpenAI(model_name="gpt-4o-mini", temperature=0.3)
        prompt = PromptTemplate(template=prompt_template, input_variables=["context", "question"])
        chain = load_qa_chain(model, chain_type="stuff", prompt=prompt)
        return chain
    except Exception as e:
        st.error(f"Lỗi khởi tạo mô hình: {e}")
        return None


# =============================
# 💬 Xử lý câu hỏi người dùng
# =============================
def user_input(user_question):
    try:
        embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
        if not os.path.exists("faiss_index"):
            st.error("❌ Chưa có dữ liệu. Hãy tải và phân tích tài liệu trước.")
            return None
        new_db = FAISS.load_local("faiss_index", embeddings, allow_dangerous_deserialization=True)
        docs = new_db.similarity_search(user_question)
        chain = get_conversation_chain()
        if not chain:
            return None

        response = chain({"input_documents": docs, "question": user_question}, return_only_outputs=True)
        answer = response["output_text"]

        if "chat_history" not in st.session_state:
            st.session_state.chat_history = []
        st.session_state.chat_history.append({"question": user_question, "answer": answer})
        return answer
    except Exception as e:
        st.error(f"Lỗi xử lý câu hỏi: {e}")
        return None


# =============================
# 📊 Thống kê câu hỏi
# =============================
def show_statistics():
    st.markdown("## 📊 Thống kê câu hỏi")

    if "chat_history" not in st.session_state or not st.session_state.chat_history:
        st.info("Chưa có dữ liệu để thống kê.")
        return

    questions = [chat["question"] for chat in st.session_state.chat_history]
    st.write(f"Tổng số câu hỏi: **{len(questions)}**")

    all_words = " ".join(questions).lower().split()
    stopwords = set(STOPWORDS) | {"câu", "hỏi", "nào", "gì", "là", "cho", "về", "có", "trong", "và", "nhưng", "thì", "với", "từ", "đến"}

    words_filtered = [w for w in all_words if w not in stopwords and len(w) > 3]
    top_words = dict(Counter(words_filtered).most_common(10))

    plt.figure(figsize=(6, 4))
    sns.barplot(x=list(top_words.values()), y=list(top_words.keys()), palette="viridis")
    plt.title("Top 10 từ khóa được dùng nhiều nhất")
    st.pyplot(plt)


# =============================
# 📄 Xuất lịch sử chat ra file Word
# =============================
def export_chat_history_to_word(chat_history):
    doc = Document()
    doc.add_heading("Lịch sử hỏi đáp ChatBot", level=1)

    for i, chat in enumerate(chat_history, 1):
        doc.add_paragraph(f"Câu hỏi {i}:", style="List Number")
        doc.add_paragraph(chat["question"])
        doc.add_paragraph(f"Câu trả lời {i}:", style="List Number")
        doc.add_paragraph(chat["answer"])
        doc.add_paragraph()

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# =============================
# 🚀 Giao diện Streamlit
# =============================
def main():
    st.set_page_config(page_title="ChatBot phân tích tài liệu (OpenAI)", page_icon="🤖")
    st.title("🤖 ChatBot phân tích tài liệu (OpenAI)")

    user_question = st.text_input("📌 Nhập câu hỏi của bạn sau khi phân tích tài liệu")

    if user_question:
        answer = user_input(user_question)
        if answer:
            st.markdown("### 💬 Trả lời:")
            st.write(answer)

    if "chat_history" in st.session_state and st.session_state.chat_history:
        st.markdown("---")
        st.markdown("## 🕑 Lịch sử chat")
        for i, chat in enumerate(st.session_state.chat_history):
            st.markdown(f"**Q{i+1}:** {chat['question']}")
            st.markdown(f"**A{i+1}:** {chat['answer']}")

        if st.button("📄 Xuất file Word"):
            word_file = export_chat_history_to_word(st.session_state.chat_history)
            st.download_button(
                label="📥 Tải file lịch sử hỏi đáp",
                data=word_file,
                file_name="chat_history.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

    if st.checkbox("📈 Hiện thống kê câu hỏi"):
        show_statistics()

    with st.sidebar:
        st.title("📁 Tải tài liệu")
        docs = st.file_uploader(
            "Tải tài liệu của bạn lên",
            accept_multiple_files=True,
            type=["pdf", "docx", "txt", "pptx", "png", "jpg", "jpeg", "bmp"]
        )

        if st.button("Phân tích tài liệu"):
            if not docs:
                st.error("Vui lòng tải tài liệu trước.")
            else:
                with st.spinner("🔍 Đang xử lý tài liệu..."):
                    raw_text = get_text(docs)
                    if raw_text:
                        chunks = get_text_chunk(raw_text)
                        if chunks:
                            get_vector_store(chunks)
                        else:
                            st.error("Không thể chia nhỏ nội dung tài liệu.")
                    else:
                        st.error("Không thể đọc nội dung tài liệu.")


if __name__ == "__main__":
    main()
